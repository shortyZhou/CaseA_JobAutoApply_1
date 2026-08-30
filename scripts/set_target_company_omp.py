from pathlib import Path
import json, shutil, datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(r'C:/Users/24779/Desktop/AI related/Case A_Job Auto Apply')
now = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')

# 1) Set Excel Sheet 1 to a single intended company: OMP
xlsx = ROOT / 'excel' / '岗位搜索与简历投递.xlsx'
wb = load_workbook(xlsx)
ws = wb['意向投递公司']
headers = [c.value for c in ws[1]]
row = {
    'company_id': 'omp',
    'Company Name': 'OMP',
    '国家/地区': '待确认',
    '城市': '待确认',
    '是否目标地区有分公司': '待确认',
    '行业': 'Supply Chain Planning Software / Consulting',
    '岗位大类': 'Technical Consultant / Analytics / Operations Research / Supply Chain',
    '匹配简历': 'profiles/profile_en.json; resumes/pdf/Resume_EN.pdf',
    'Career 官网链接': 'https://careers.omp.com/',
    'ATS 类型': '自建 / 未知',
    'ATS Board 链接': '',
    'LinkedIn 公司岗位链接': 'https://www.linkedin.com/company/omp-supply-chain-planning/jobs',
    '信息来源': '用户指定 / Agent 搜索补全',
    '最近检查时间': now,
    '备注': 'autoapply-mcp 未发现 Greenhouse / Lever / Ashby board；后续优先用 Hermes browser automation 检查 careers.omp.com。'
}
# Preserve header; clear template/data rows because user explicitly wants only one intended company.
if ws.max_row > 1:
    ws.delete_rows(2, ws.max_row - 1)
ws.append([row.get(h, '') for h in headers])
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='1F4E78')
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
wb.save(xlsx)

# 2) Make autoapply board config reflect that no supported ATS board is currently configured.
# OMP appears to use a custom careers site, so putting fake Greenhouse/Lever/Ashby data here would be wrong.
companies_path = ROOT / '.autoapply' / 'companies.json'
backup_path = ROOT / '.autoapply' / f'companies.backup-before-omp-{datetime.datetime.now().strftime("%Y%m%d_%H%M%S")}.json'
if companies_path.exists():
    old = json.loads(companies_path.read_text(encoding='utf-8'))
    if old.get('companies') != []:
        shutil.copy2(companies_path, backup_path)
        backup = str(backup_path)
    else:
        backup = None
else:
    old = {'version': 1, 'companies': []}
    backup = None
companies_path.write_text(json.dumps({'version': 1, 'companies': []}, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

# 3) Update preferences with single target company without overwriting other preferences.
pref_path = ROOT / 'profiles' / 'application_preferences.json'
prefs = json.loads(pref_path.read_text(encoding='utf-8')) if pref_path.exists() else {'version': 1}
prefs['target_companies'] = ['OMP']
prefs['target_company_notes'] = {
    'OMP': {
        'career_url': 'https://careers.omp.com/',
        'linkedin_jobs_url': 'https://www.linkedin.com/company/omp-supply-chain-planning/jobs',
        'ats_status': 'No Greenhouse / Lever / Ashby board resolved; likely custom careers site.',
        'recommended_tool': 'Hermes browser automation only unless a supported ATS board is later identified.'
    }
}
pref_path.write_text(json.dumps(prefs, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

# 4) Audit log
log_path = ROOT / 'logs' / 'setup_runs.md'
with log_path.open('a', encoding='utf-8') as f:
    f.write(f"\n## {now} Set intended company to OMP\n\n")
    f.write("- Updated Excel sheet `意向投递公司` to one row: OMP.\n")
    f.write("- OMP careers URL: https://careers.omp.com/\n")
    f.write("- `autoapply-mcp` resolver found no Greenhouse / Lever / Ashby board for OMP.\n")
    f.write("- Replaced `.autoapply/companies.json` with an empty supported-ATS board list; old list backed up if non-empty.\n")
    if backup:
        f.write(f"- Backup: `{backup}`\n")

print(json.dumps({
    'excel_updated': str(xlsx),
    'sheet': '意向投递公司',
    'company': row,
    'autoapply_companies': str(companies_path),
    'autoapply_supported_boards_count': 0,
    'backup': backup,
    'preferences_updated': str(pref_path)
}, ensure_ascii=False, indent=2))
