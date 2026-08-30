from pathlib import Path
import json, shutil, datetime, hashlib, re
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(r'C:/Users/24779/Desktop/AI related/Case A_Job Auto Apply')
NOW = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
TODAY = datetime.datetime.now().strftime('%Y-%m-%d')
XLSX = ROOT / 'excel' / '岗位搜索与简历投递.xlsx'

JOBS = [
    {
        'job_id': 'omp-4898879101',
        'company_id': 'omp',
        'Company Name': 'OMP',
        'Job Title': 'Integration Consultant',
        'Role Category': 'Technical Consultant / Integration / Supply Chain Planning',
        'Country/Region': 'China / APAC',
        'City': 'Shanghai',
        'Work Mode': 'hybrid',
        'Language': 'English',
        'JD Summary': 'Implement OMP integration solutions into customer enterprise systems; design workflows, write integration scripts, data mapping/modeling/loading, specifications, testing and customer-needs analysis.',
        'Key Requirements': "Master's in CS/Mathematics/Industrial Engineering or similar; solid IT background; 1-3 years integration/troubleshooting/support; analytical/problem-solving; English and Chinese; travel. Bonus: SAP, APIs, SQL/R/JavaScript/Python, ETL.",
        'JD URL': 'https://careers.omp.com/jobs/china-shanghai/integration-consultant-4898879101',
        'Application URL': 'https://careers.omp.com/jobs/china-shanghai/integration-consultant-4898879101',
        'Source Platform': 'Company Site / OMP careers',
        'Match Score': 86,
        'Match Rationale': 'Strong fit with NUS Industrial & Systems Engineering, KPMG technical consulting internship, Python/SQL/data mapping, analytical and bilingual English/Chinese profile.',
        'Risk / Gaps': 'Role asks for 1-3 years integration implementation/support and SAP/enterprise-system exposure; travel and China work authorization need user confirmation.',
        'Discovered At': NOW,
        'Last Checked At': NOW,
        'Is Applied': 'No',
        'User Decision': 'Pending',
        'Notes': 'APAC target; official OMP careers page; do not auto-submit.'
    },
    {
        'job_id': 'omp-4950152101',
        'company_id': 'omp',
        'Company Name': 'OMP',
        'Job Title': 'Product Consultant',
        'Role Category': 'Product Consultant / Analytics / Supply Chain Planning',
        'Country/Region': 'China / APAC',
        'City': 'Shanghai',
        'Work Mode': 'hybrid',
        'Language': 'English',
        'JD Summary': 'Support customer project teams in supply-chain planning product domains; help solution architects and implementation consultants; solve customer-specific product challenges; brief product managers on product gaps and differentiators; build prototypes for customer problems.',
        'Key Requirements': "Master's/PhD in Engineering, Business Administration, Applied Economics, Mathematics, CS or similar; excellent English and Chinese; analytical/problem-solving; IT affinity; basic programming; travel. Bonus: supply-chain planning/business software, mathematical modeling or ML.",
        'JD URL': 'https://careers.omp.com/jobs/china-shanghai/product-consultant-4950152101',
        'Application URL': 'https://careers.omp.com/jobs/china-shanghai/product-consultant-4950152101',
        'Source Platform': 'Company Site / OMP careers',
        'Match Score': 91,
        'Match Rationale': 'Best fit among OMP APAC roles: Industrial & Systems Engineering, Operations Research, mathematical modeling, ML, SQL/basic programming, consulting/data-analysis background, English/Chinese requirement.',
        'Risk / Gaps': 'Supply-chain planning or business-software implementation experience is bonus but not clearly shown as core work experience; travel and China work authorization need confirmation.',
        'Discovered At': NOW,
        'Last Checked At': NOW,
        'Is Applied': 'No',
        'User Decision': 'Pending',
        'Notes': 'Recommended first pilot role for semi-automated application demo.'
    },
    {
        'job_id': 'omp-4669760101',
        'company_id': 'omp',
        'Company Name': 'OMP',
        'Job Title': 'Supply Chain (Implementation) Consultant',
        'Role Category': 'Supply Chain Planning / Implementation Consultant',
        'Country/Region': 'China / APAC',
        'City': 'Shanghai',
        'Work Mode': 'hybrid',
        'Language': 'English',
        'JD Summary': 'Analyze supply-chain planning challenges, assist solution design, implement and validate OMP solutions, support go-live, train customers, and document technical/functional solution notes.',
        'Key Requirements': "Strong analytical skills; preferably Master's specializing in Production Management, Logistics or similar; strong IT affinity; basic programming; travel; fluent English; Chinese required.",
        'JD URL': 'https://careers.omp.com/jobs/china-shanghai/supply-chain-implementation-consultant-4669760101',
        'Application URL': 'https://careers.omp.com/jobs/china-shanghai/supply-chain-implementation-consultant-4669760101',
        'Source Platform': 'Company Site / OMP careers',
        'Match Score': 88,
        'Match Rationale': 'Good fit with Industrial & Systems Engineering, Operations Research, forecasting/modeling, technical consulting internship, Python and analytical skills; English/Chinese requirement likely aligned.',
        'Risk / Gaps': 'Direct supply-chain implementation experience not obvious from current profile; travel and China work authorization need confirmation.',
        'Discovered At': NOW,
        'Last Checked At': NOW,
        'Is Applied': 'No',
        'User Decision': 'Pending',
        'Notes': 'Strong demo candidate; official OMP careers page; do not auto-submit.'
    },
    {
        'job_id': 'omp-4740094101',
        'company_id': 'omp',
        'Company Name': 'OMP',
        'Job Title': 'Supply Chain (Implementation) Consultant',
        'Role Category': 'Supply Chain Planning / Implementation Consultant',
        'Country/Region': 'Australia / APAC',
        'City': 'Melbourne',
        'Work Mode': 'hybrid',
        'Language': 'English',
        'JD Summary': 'Analyze supply-chain planning challenges, design and implement OMP solutions with customer project teams, validate go-live, train customers, and document implemented solutions.',
        'Key Requirements': "Strong analytical skills; preferably Master's specializing in Production Management, Logistics or similar; IT affinity; basic programming; travel; fluent English.",
        'JD URL': 'https://careers.omp.com/jobs/australia-melbourne/supply-chain-implementation-consultant-4740094101',
        'Application URL': 'https://careers.omp.com/jobs/australia-melbourne/supply-chain-implementation-consultant-4740094101',
        'Source Platform': 'Company Site / OMP careers',
        'Match Score': 76,
        'Match Rationale': 'Technical and analytical fit is reasonable; supply-chain/OR background aligns with the role theme.',
        'Risk / Gaps': 'Australia work authorization, relocation/location, and direct supply-chain implementation experience need confirmation.',
        'Discovered At': NOW,
        'Last Checked At': NOW,
        'Is Applied': 'No',
        'User Decision': 'Pending',
        'Notes': 'Lower priority unless user wants Australia roles.'
    },
    {
        'job_id': 'omp-4950311101',
        'company_id': 'omp',
        'Company Name': 'OMP',
        'Job Title': 'Supply Chain (Implementation) Consultant',
        'Role Category': 'Supply Chain Planning / Implementation Consultant',
        'Country/Region': 'Japan / APAC',
        'City': 'Tokyo',
        'Work Mode': 'hybrid',
        'Language': 'English',
        'JD Summary': 'Analyze supply-chain planning challenges, design and implement OMP planning solutions, validate go-live, train customers, and document technical/functional solution notes.',
        'Key Requirements': "Strong analytical skills; preferably Master's in Production Management, Logistics or similar; IT affinity; basic programming; travel; fluent English and Japanese.",
        'JD URL': 'https://careers.omp.com/jobs/japan-tokyo/supply-chain-implementation-consultant--4950311101',
        'Application URL': 'https://careers.omp.com/jobs/japan-tokyo/supply-chain-implementation-consultant--4950311101',
        'Source Platform': 'Company Site / OMP careers',
        'Match Score': 62,
        'Match Rationale': 'Role theme aligns with Industrial & Systems Engineering and OR, but language/location constraints reduce fit.',
        'Risk / Gaps': 'Japanese fluency required; Japan work authorization/relocation need confirmation. Keep as low-priority APAC reference.',
        'Discovered At': NOW,
        'Last Checked At': NOW,
        'Is Applied': 'No',
        'User Decision': 'Pending',
        'Notes': 'Not recommended unless Japanese requirement is satisfied.'
    }
]

# Update preferences and mapping.
pref_path = ROOT / 'profiles' / 'application_preferences.json'
prefs = json.loads(pref_path.read_text(encoding='utf-8'))
prefs['target_regions'] = ['APAC']
prefs['target_industries'] = ['Supply Chain Planning Software', 'Consulting', 'Operations Research / Optimization', 'Enterprise Software']
prefs['target_role_categories'] = ['Supply Chain Planning', 'Technical Consultant', 'Integration Consultant', 'Product Consultant', 'Operations Research']
prefs['target_companies'] = ['OMP']
prefs['notes'] = 'User confirmed APAC and OMP Supply Chain Planning / Technical Consultant / Operations Research as the test scope.'
pref_path.write_text(json.dumps(prefs, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

map_path = ROOT / 'profiles' / 'resume_role_mapping.json'
mapping = json.loads(map_path.read_text(encoding='utf-8')) if map_path.exists() else {'version': 1, 'mappings': []}
mapping['mappings'] = [{
    'profile': 'profile_en.json',
    'resume_file': 'resumes/pdf/Resume_EN.pdf',
    'target_regions': ['APAC'],
    'target_companies': ['OMP'],
    'industries': ['Supply Chain Planning Software', 'Consulting', 'Operations Research / Optimization', 'Enterprise Software'],
    'role_categories': ['Supply Chain Planning', 'Technical Consultant', 'Integration Consultant', 'Product Consultant', 'Operations Research'],
    'notes': 'Updated after user confirmation for OMP APAC demo scope.'
}]
map_path.write_text(json.dumps(mapping, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

# Update workbook with backups.
backup = XLSX.with_name(f'{XLSX.stem}.backup-before-omp-apac-{datetime.datetime.now().strftime("%Y%m%d_%H%M%S")}{XLSX.suffix}')
shutil.copy2(XLSX, backup)
wb = load_workbook(XLSX)

# Sheet 1: update OMP target row region/direction.
ws_company = wb['意向投递公司']
company_headers = [c.value for c in ws_company[1]]
company_col = {h:i+1 for i,h in enumerate(company_headers)}
# Ensure one OMP row exists.
if ws_company.max_row < 2:
    ws_company.append(['' for _ in company_headers])
# Clear any other data rows; user explicitly chose one intended company.
if ws_company.max_row > 2:
    ws_company.delete_rows(3, ws_company.max_row - 2)
company_updates = {
    'company_id': 'omp',
    'Company Name': 'OMP',
    '国家/地区': 'APAC',
    '城市': 'Shanghai / Melbourne / Tokyo',
    '是否目标地区有分公司': '是',
    '行业': 'Supply Chain Planning Software / Consulting',
    '岗位大类': 'Supply Chain Planning / Technical Consultant / Operations Research',
    '匹配简历': 'profiles/profile_en.json; resumes/pdf/Resume_EN.pdf',
    'Career 官网链接': 'https://careers.omp.com/',
    'ATS 类型': '自建 / 未知',
    'ATS Board 链接': '',
    'LinkedIn 公司岗位链接': 'https://www.linkedin.com/company/omp-supply-chain-planning/jobs',
    '信息来源': '用户指定 / Agent 搜索补全',
    '最近检查时间': NOW,
    '备注': 'User confirmed APAC and OMP Supply Chain Planning / Technical Consultant / Operations Research demo scope. autoapply-mcp did not find Greenhouse/Lever/Ashby board; use Hermes browser automation.'
}
for h,v in company_updates.items():
    if h in company_col:
        ws_company.cell(row=2, column=company_col[h], value=v)

# Sheet 2: upsert jobs by job_id or JD URL.
ws_jobs = wb['开放岗位列表']
headers = [c.value for c in ws_jobs[1]]
col = {h:i+1 for i,h in enumerate(headers)}
existing_by_id = {}
existing_by_url = {}
for r in range(2, ws_jobs.max_row + 1):
    jid = ws_jobs.cell(r, col['job_id']).value if 'job_id' in col else None
    url = ws_jobs.cell(r, col['JD URL']).value if 'JD URL' in col else None
    if jid: existing_by_id[str(jid)] = r
    if url: existing_by_url[str(url)] = r
inserted = 0
updated = 0
for job in JOBS:
    rownum = existing_by_id.get(job['job_id']) or existing_by_url.get(job['JD URL'])
    if rownum:
        # Only update machine-owned check/time + blank fields; don't overwrite user decision/applied/notes if manually changed.
        for h,v in job.items():
            if h not in col:
                continue
            cell = ws_jobs.cell(rownum, col[h])
            if h in ('Last Checked At',):
                cell.value = v
            elif cell.value in (None, ''):
                cell.value = v
        updated += 1
    else:
        ws_jobs.append([job.get(h, '') for h in headers])
        inserted += 1

# Style header and freeze panes.
for ws in (ws_company, ws_jobs):
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1F4E78')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
# Reasonable widths for job sheet.
widths = {'A':18,'B':12,'C':16,'D':36,'E':34,'F':18,'G':18,'H':12,'I':12,'J':60,'K':65,'L':70,'M':70,'N':24,'O':12,'P':60,'Q':55,'R':20,'S':20,'T':12,'U':16,'V':36}
for letter,width in widths.items():
    ws_jobs.column_dimensions[letter].width = width

wb.save(XLSX)

# Audit log.
log_path = ROOT / 'logs' / 'job_search_runs.md'
with log_path.open('a', encoding='utf-8') as f:
    f.write(f'\n## {NOW} OMP APAC job search\n\n')
    f.write('- User confirmed target region: APAC.\n')
    f.write('- User accepted OMP Supply Chain Planning / Technical Consultant / Operations Research directions.\n')
    f.write('- Source: OMP official careers site and individual JD pages.\n')
    f.write(f'- Excel updated: `{XLSX}`\n')
    f.write(f'- Backup before write: `{backup}`\n')
    f.write(f'- Inserted rows: {inserted}; updated existing rows: {updated}.\n')
    f.write('- Safety: no application submitted; rows remain `User Decision = Pending`, `Is Applied = No`.\n\n')
    for job in JOBS:
        f.write(f"- {job['Job Title']} — {job['City']} — {job['JD URL']} — score {job['Match Score']}\n")

print(json.dumps({'excel': str(XLSX), 'backup': str(backup), 'inserted': inserted, 'updated': updated, 'jobs': JOBS}, ensure_ascii=False, indent=2))
