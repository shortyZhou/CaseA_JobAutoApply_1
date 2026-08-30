"""Maintain the Case A master job-application Excel workbook.

Safety rules:
- Writes are incremental and create a timestamped backup before every save.
- Existing user-entered cells are not overwritten unless --overwrite is passed.
- Final application submission is never assumed; use mark-applied only after
  explicit user direction, and use --submitted No for filled-not-submitted cases.

Examples:
  python scripts/update_job_excel.py summary
  python scripts/update_job_excel.py append-job --json '{"job_id":"demo","Company Name":"DemoCo","Job Title":"Analyst"}'
  python scripts/update_job_excel.py update-decision --job-id demo --decision Apply
  python scripts/update_job_excel.py mark-applied --job-id demo --submitted Yes --status submitted
  python scripts/update_job_excel.py update-status --job-id demo --status "under review"
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import shutil
import sys
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "excel" / "岗位搜索与简历投递.xlsx"
BACKUP_DIR = ROOT / "backups" / "excel"
LOG_PATH = ROOT / "logs" / "job_search_runs.md"

OPEN_JOBS_SHEET = "开放岗位列表"
APPLIED_SHEET = "已投递记录"
COMPANIES_SHEET = "意向投递公司"

OPEN_JOBS_HEADERS = [
    "job_id",
    "company_id",
    "Company Name",
    "Job Title",
    "Role Category",
    "Country/Region",
    "City",
    "Work Mode",
    "Language",
    "JD Summary",
    "Key Requirements",
    "JD URL",
    "Application URL",
    "Source Platform",
    "Match Score",
    "Match Rationale",
    "Risk / Gaps",
    "Discovered At",
    "Last Checked At",
    "Is Applied",
    "User Decision",
    "Notes",
]

APPLIED_HEADERS = [
    "application_id",
    "job_id",
    "Company Name",
    "Job Title",
    "Country/Region",
    "City",
    "Resume/Profile Used",
    "JD URL",
    "Application URL",
    "Platform",
    "Applied At",
    "Tool Combination",
    "User Submitted Final Application",
    "Current Status",
    "Status Updated At",
    "Status 1",
    "Updated At 1",
    "Status 2",
    "Updated At 2",
    "Status 3",
    "Updated At 3",
    "Status 4",
    "Updated At 4",
    "Notes",
]

COMPANY_HEADERS = [
    "company_id",
    "Company Name",
    "国家/地区",
    "城市",
    "是否目标地区有分公司",
    "行业",
    "岗位大类",
    "匹配简历",
    "Career 官网链接",
    "ATS 类型",
    "ATS Board 链接",
    "LinkedIn 公司岗位链接",
    "信息来源",
    "最近检查时间",
    "备注",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def now_str() -> str:
    return dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def slug(value: str) -> str:
    raw = (value or "").strip().lower()
    keep = [ch if ch.isalnum() else "-" for ch in raw]
    return "-".join("".join(keep).split("-"))[:60] or "unknown"


def short_hash(value: str) -> str:
    return hashlib.sha1(value.encode("utf-8")).hexdigest()[:10]


def normalize_blank(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def read_payload(args: argparse.Namespace) -> Any:
    if getattr(args, "json_text", None):
        return json.loads(args.json_text)
    if getattr(args, "data", None):
        return json.loads(Path(args.data).read_text(encoding="utf-8"))
    return {}


def row_to_dict(ws, row_index: int) -> dict[str, Any]:
    headers = get_headers(ws)
    return {h: ws.cell(row_index, idx).value for h, idx in headers.items()}


def get_headers(ws) -> dict[str, int]:
    headers: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(1, col).value
        if value:
            headers[str(value).strip()] = col
    return headers


def style_sheet(ws) -> None:
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = min(max(len(str(ws.cell(1, col).value or "")) + 4, 14), 42)


def ensure_sheet(wb, name: str, headers: list[str]):
    ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
    existing = get_headers(ws)
    if ws.max_row == 1 and not any(ws.cell(1, c).value for c in range(1, ws.max_column + 1)):
        existing = {}
    for header in headers:
        if header not in existing:
            ws.cell(1, ws.max_column + 1 if existing else len(existing) + 1, header)
            existing = get_headers(ws)
    # Keep canonical header order for empty sheets only. Non-empty sheets preserve user order.
    if ws.max_row <= 1:
        for col, header in enumerate(headers, start=1):
            ws.cell(1, col, header)
    style_sheet(ws)
    return ws


def load_or_create_workbook(path: Path):
    if path.exists():
        wb = load_workbook(path)
    else:
        path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        wb.active.title = COMPANIES_SHEET
    ensure_sheet(wb, COMPANIES_SHEET, COMPANY_HEADERS)
    ensure_sheet(wb, OPEN_JOBS_SHEET, OPEN_JOBS_HEADERS)
    ensure_sheet(wb, APPLIED_SHEET, APPLIED_HEADERS)
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 3:
        # Remove only the default empty sheet, never user-created sheets with content.
        ws = wb["Sheet"]
        if ws.max_row == 1 and ws.max_column == 1 and ws["A1"].value is None:
            wb.remove(ws)
    return wb


def backup_workbook(path: Path, action: str) -> str | None:
    if not path.exists():
        return None
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    target = BACKUP_DIR / f"{path.stem}.{action}.{stamp}{path.suffix}"
    shutil.copy2(path, target)
    return str(target)


def save_with_backup(wb, path: Path, action: str) -> str | None:
    backup = backup_workbook(path, action)
    for ws in wb.worksheets:
        style_sheet(ws)
    wb.save(path)
    return backup


def append_log(action: str, result: dict[str, Any]) -> None:
    LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    with LOG_PATH.open("a", encoding="utf-8") as f:
        f.write(f"\n## {now_str()} update_job_excel: {action}\n\n")
        f.write("```json\n")
        f.write(json.dumps(result, ensure_ascii=False, indent=2, default=str))
        f.write("\n```\n")


def find_row(ws, *, job_id: str | None = None, application_id: str | None = None, urls: Iterable[str] = ()) -> int | None:
    headers = get_headers(ws)
    url_set = {u.strip() for u in urls if u and str(u).strip()}
    for r in range(2, ws.max_row + 1):
        if job_id and "job_id" in headers and normalize_blank(ws.cell(r, headers["job_id"]).value) == job_id:
            return r
        if application_id and "application_id" in headers and normalize_blank(ws.cell(r, headers["application_id"]).value) == application_id:
            return r
        for h in ("JD URL", "Application URL"):
            if h in headers and normalize_blank(ws.cell(r, headers[h]).value) in url_set:
                return r
    return None


def set_cell(ws, row: int, header: str, value: Any, *, overwrite: bool = False) -> bool:
    if value is None or header not in get_headers(ws):
        return False
    headers = get_headers(ws)
    cell = ws.cell(row, headers[header])
    if overwrite or normalize_blank(cell.value) == "":
        old = cell.value
        cell.value = value
        return old != value
    return False


def append_jobs(args: argparse.Namespace) -> dict[str, Any]:
    payload = read_payload(args)
    jobs = payload if isinstance(payload, list) else [payload]
    wb = load_or_create_workbook(Path(args.workbook))
    ws = wb[OPEN_JOBS_SHEET]
    added: list[str] = []
    updated: list[str] = []
    skipped: list[str] = []
    timestamp = now_str()

    for job in jobs:
        if not isinstance(job, dict):
            raise ValueError("Job payload must be a JSON object or list of objects")
        job = {str(k): v for k, v in job.items()}
        if not normalize_blank(job.get("job_id")):
            seed = normalize_blank(job.get("Application URL")) or normalize_blank(job.get("JD URL")) or f"{job.get('Company Name','')}-{job.get('Job Title','')}"
            job["job_id"] = f"job-{slug(str(job.get('Company Name', 'company')))}-{short_hash(seed)}"
        job.setdefault("Language", "English")
        job.setdefault("Discovered At", timestamp)
        job.setdefault("Last Checked At", timestamp)
        job.setdefault("Is Applied", "No")
        job.setdefault("User Decision", "Pending")

        row = find_row(ws, job_id=normalize_blank(job["job_id"]), urls=[job.get("JD URL", ""), job.get("Application URL", "")])
        if row is None:
            ws.append([job.get(h, "") for h in OPEN_JOBS_HEADERS])
            added.append(str(job["job_id"]))
        else:
            changed = False
            # Safe upsert: only fill blanks unless --overwrite; Last Checked At is a system-maintained field.
            for header in OPEN_JOBS_HEADERS:
                ow = bool(args.overwrite) or header == "Last Checked At"
                changed = set_cell(ws, row, header, job.get(header), overwrite=ow) or changed
            (updated if changed else skipped).append(str(job["job_id"]))

    backup = save_with_backup(wb, Path(args.workbook), "append-job") if added or updated else None
    result = {"action": "append-job", "added": added, "updated": updated, "skipped_no_change": skipped, "backup": backup, "workbook": str(args.workbook)}
    append_log("append-job", result)
    return result


def update_decision(args: argparse.Namespace) -> dict[str, Any]:
    if args.decision not in {"Apply", "Skip", "Pending"}:
        raise ValueError("--decision must be Apply, Skip, or Pending")
    wb = load_or_create_workbook(Path(args.workbook))
    ws = wb[OPEN_JOBS_SHEET]
    row = find_row(ws, job_id=args.job_id)
    if row is None:
        raise KeyError(f"job_id not found in {OPEN_JOBS_SHEET}: {args.job_id}")
    changed = set_cell(ws, row, "User Decision", args.decision, overwrite=True)
    if args.notes:
        headers = get_headers(ws)
        cell = ws.cell(row, headers["Notes"])
        existing = normalize_blank(cell.value)
        note = f"[{now_str()}] {args.notes}"
        cell.value = f"{existing}\n{note}" if existing else note
        changed = True
    backup = save_with_backup(wb, Path(args.workbook), "update-decision") if changed else None
    result = {"action": "update-decision", "job_id": args.job_id, "decision": args.decision, "changed": changed, "backup": backup}
    append_log("update-decision", result)
    return result


def application_id_for(job_id: str) -> str:
    return f"app-{slug(job_id)}"


def mark_applied(args: argparse.Namespace) -> dict[str, Any]:
    if args.submitted not in {"Yes", "No", "Pending"}:
        raise ValueError("--submitted must be Yes, No, or Pending")
    wb = load_or_create_workbook(Path(args.workbook))
    jobs_ws = wb[OPEN_JOBS_SHEET]
    app_ws = wb[APPLIED_SHEET]
    job_row_index = find_row(jobs_ws, job_id=args.job_id)
    if job_row_index is None:
        raise KeyError(f"job_id not found in {OPEN_JOBS_SHEET}: {args.job_id}")
    job = row_to_dict(jobs_ws, job_row_index)
    app_id = args.application_id or application_id_for(args.job_id)
    app_row_index = find_row(app_ws, job_id=args.job_id, application_id=app_id)
    timestamp = args.applied_at or now_str()
    status = args.status or ("submitted" if args.submitted == "Yes" else "filled_not_submitted")

    app_data = {
        "application_id": app_id,
        "job_id": args.job_id,
        "Company Name": job.get("Company Name"),
        "Job Title": job.get("Job Title"),
        "Country/Region": job.get("Country/Region"),
        "City": job.get("City"),
        "Resume/Profile Used": args.resume_profile,
        "JD URL": job.get("JD URL"),
        "Application URL": job.get("Application URL"),
        "Platform": job.get("Source Platform") or args.platform,
        "Applied At": timestamp if args.submitted == "Yes" else "",
        "Tool Combination": args.tool_combination,
        "User Submitted Final Application": args.submitted,
        "Current Status": status,
        "Status Updated At": timestamp,
        "Status 1": status,
        "Updated At 1": timestamp,
        "Notes": args.notes or ("Filled to final review step; not submitted." if args.submitted == "No" else ""),
    }

    created = False
    changed = False
    if app_row_index is None:
        app_ws.append([app_data.get(h, "") for h in get_headers(app_ws)])
        created = True
        changed = True
    else:
        for header, value in app_data.items():
            # Existing application records may contain user notes/status history; only overwrite explicit top-level fields.
            ow = header in {"User Submitted Final Application", "Current Status", "Status Updated At"}
            changed = set_cell(app_ws, app_row_index, header, value, overwrite=ow) or changed

    # Record Is Applied in the open-jobs sheet without claiming submission when not submitted.
    is_applied_value = "Yes" if args.submitted == "Yes" else "No" if args.submitted == "No" else "Pending"
    changed = set_cell(jobs_ws, job_row_index, "Is Applied", is_applied_value, overwrite=True) or changed
    if args.submitted == "Yes":
        changed = set_cell(jobs_ws, job_row_index, "User Decision", "Apply", overwrite=True) or changed

    backup = save_with_backup(wb, Path(args.workbook), "mark-applied") if changed else None
    result = {"action": "mark-applied", "application_id": app_id, "job_id": args.job_id, "created": created, "changed": changed, "submitted": args.submitted, "status": status, "backup": backup}
    append_log("mark-applied", result)
    return result


def next_status_slot(ws, row: int) -> tuple[str, str]:
    headers = get_headers(ws)
    n = 1
    while True:
        status_h = f"Status {n}"
        time_h = f"Updated At {n}"
        if status_h not in headers:
            notes_col = headers.get("Notes", ws.max_column + 1)
            ws.insert_cols(notes_col, 2)
            ws.cell(1, notes_col, status_h)
            ws.cell(1, notes_col + 1, time_h)
            return status_h, time_h
        if normalize_blank(ws.cell(row, headers[status_h]).value) == "":
            return status_h, time_h
        n += 1


def update_status(args: argparse.Namespace) -> dict[str, Any]:
    wb = load_or_create_workbook(Path(args.workbook))
    ws = wb[APPLIED_SHEET]
    row = find_row(ws, job_id=args.job_id, application_id=args.application_id)
    if row is None:
        raise KeyError("application not found; pass --application-id or --job-id for an existing applied record")
    headers = get_headers(ws)
    latest = normalize_blank(ws.cell(row, headers["Current Status"]).value)
    if latest == args.status:
        result = {"action": "update-status", "changed": False, "reason": "same_status", "status": args.status}
        append_log("update-status", result)
        return result
    timestamp = args.updated_at or now_str()
    status_h, time_h = next_status_slot(ws, row)
    set_cell(ws, row, status_h, args.status, overwrite=True)
    set_cell(ws, row, time_h, timestamp, overwrite=True)
    set_cell(ws, row, "Current Status", args.status, overwrite=True)
    set_cell(ws, row, "Status Updated At", timestamp, overwrite=True)
    if args.notes:
        headers = get_headers(ws)
        cell = ws.cell(row, headers["Notes"])
        existing = normalize_blank(cell.value)
        note = f"[{timestamp}] {args.notes}"
        cell.value = f"{existing}\n{note}" if existing else note
    backup = save_with_backup(wb, Path(args.workbook), "update-status")
    result = {"action": "update-status", "changed": True, "previous_status": latest, "new_status": args.status, "slot": status_h, "backup": backup}
    append_log("update-status", result)
    return result


def summary(args: argparse.Namespace) -> dict[str, Any]:
    wb = load_or_create_workbook(Path(args.workbook))
    data: dict[str, Any] = {"workbook": str(args.workbook), "sheets": {}}
    for sheet in [COMPANIES_SHEET, OPEN_JOBS_SHEET, APPLIED_SHEET]:
        ws = wb[sheet]
        non_empty_rows = 0
        for r in range(2, ws.max_row + 1):
            if any(normalize_blank(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)):
                non_empty_rows += 1
        data["sheets"][sheet] = {"rows": non_empty_rows, "columns": list(get_headers(ws).keys())}
    return data


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Incrementally maintain the Case A application tracker workbook.")
    parser.add_argument("--workbook", default=str(DEFAULT_WORKBOOK), help="Path to the master .xlsx workbook")
    sub = parser.add_subparsers(dest="command", required=True)

    p = sub.add_parser("summary", help="Print workbook sheet/column summary")
    p.set_defaults(func=summary)

    p = sub.add_parser("append-job", help="Append or safely upsert one or more rows into 开放岗位列表")
    p.add_argument("--json", dest="json_text", help="Job JSON object/list. Prefer --data for long payloads.")
    p.add_argument("--data", help="Path to JSON object/list file")
    p.add_argument("--overwrite", action="store_true", help="Overwrite existing non-empty cells; off by default")
    p.set_defaults(func=append_jobs)

    p = sub.add_parser("update-decision", help="Set User Decision for an open job")
    p.add_argument("--job-id", required=True)
    p.add_argument("--decision", required=True, choices=["Apply", "Skip", "Pending"])
    p.add_argument("--notes")
    p.set_defaults(func=update_decision)

    p = sub.add_parser("mark-applied", help="Create/update an application record and update Is Applied")
    p.add_argument("--job-id", required=True)
    p.add_argument("--application-id")
    p.add_argument("--submitted", required=True, choices=["Yes", "No", "Pending"], help="Whether the user manually submitted the final application")
    p.add_argument("--status", help="Default: submitted if Yes; filled_not_submitted if No")
    p.add_argument("--applied-at")
    p.add_argument("--resume-profile", default="profiles/profile_en.json; resumes/pdf/Resume_EN.pdf")
    p.add_argument("--platform", default="Company Site")
    p.add_argument("--tool-combination", default="Hermes browser automation")
    p.add_argument("--notes")
    p.set_defaults(func=mark_applied)

    p = sub.add_parser("update-status", help="Append a new status change to 已投递记录")
    p.add_argument("--application-id")
    p.add_argument("--job-id")
    p.add_argument("--status", required=True)
    p.add_argument("--updated-at")
    p.add_argument("--notes")
    p.set_defaults(func=update_status)
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        result = args.func(args)
        print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
        return 0
    except Exception as exc:  # Keep CLI failures explicit for logs and terminal output.
        print(json.dumps({"error": type(exc).__name__, "message": str(exc)}, ensure_ascii=False, indent=2), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
