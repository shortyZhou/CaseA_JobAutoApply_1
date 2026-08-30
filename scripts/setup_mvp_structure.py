from pathlib import Path
import shutil, json, hashlib, datetime

root = Path(r'C:/Users/24779/Desktop/AI related/Case A_Job Auto Apply')
resume_src = root / 'Resume.pdf'
if not resume_src.exists():
    raise SystemExit('Resume.pdf not found')

def sha(p: Path) -> str:
    h = hashlib.sha256()
    with open(p, 'rb') as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()

def ensure_dir(rel: str) -> Path:
    p = root / rel
    p.mkdir(parents=True, exist_ok=True)
    return p

def copy_no_loss(src: Path, dst: Path) -> dict:
    dst.parent.mkdir(parents=True, exist_ok=True)
    if dst.exists():
        if sha(src) == sha(dst):
            return {'path': str(dst), 'action': 'exists_same'}
        alt = dst.with_name(dst.stem + '_from_project_root' + dst.suffix)
        i = 2
        while alt.exists() and sha(src) != sha(alt):
            alt = dst.with_name(f'{dst.stem}_from_project_root_{i}{dst.suffix}')
            i += 1
        if alt.exists() and sha(src) == sha(alt):
            return {'path': str(alt), 'action': 'exists_same_alt_due_conflict', 'original_conflict': str(dst)}
        shutil.copy2(src, alt)
        return {'path': str(alt), 'action': 'copied_alt_due_conflict', 'original_conflict': str(dst)}
    shutil.copy2(src, dst)
    return {'path': str(dst), 'action': 'copied'}

for rel in [
    'resumes/original', 'resumes/pdf', 'profiles', 'excel',
    'applications/generated_answers', 'applications/application_packets', 'applications/screenshots',
    'logs', 'mcp/autoapply-mcp', '.autoapply/resumes', '.autoapply/data', '.autoapply/artifacts', 'scripts'
]:
    ensure_dir(rel)

copy_results = [
    copy_no_loss(resume_src, root / 'resumes/original/Resume.pdf'),
    copy_no_loss(resume_src, root / 'resumes/pdf/Resume_EN.pdf'),
    copy_no_loss(resume_src, root / '.autoapply/resumes/Resume_EN.pdf'),
]

if (root / 'resume_data.json').exists() and not (root / 'profiles/resume_data.json').exists():
    shutil.copy2(root / 'resume_data.json', root / 'profiles/resume_data.json')
    resume_data_action = 'copied_to_profiles'
elif (root / 'profiles/resume_data.json').exists():
    resume_data_action = 'profiles_resume_data_exists'
else:
    resume_data_action = 'missing_source'

created = []

def write_json_if_missing(rel: str, data: dict) -> str:
    p = root / rel
    if not p.exists():
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        created.append(rel)
        return 'created'
    return 'exists'

write_json_if_missing('profiles/application_preferences.json', {
    'version': 1,
    'language_scope': 'English only',
    'target_regions': ['Singapore', 'Hong Kong', 'APAC', 'Remote'],
    'target_industries': ['FinTech', 'Consulting', 'Internet', 'Data Services'],
    'target_role_categories': [
        'Data Analyst', 'Business Analyst', 'Data Scientist', 'Machine Learning / AI',
        'Quantitative / FinTech', 'Technical Consultant', 'Analytics / Strategy / Operations Analytics'
    ],
    'excluded_companies': [],
    'excluded_industries': [],
    'autofill_policy': {
        'mode': 'semi_automated_only',
        'never_auto_submit': True,
        'sensitive_fields_require_user_confirmation': [
            'work authorization', 'visa status', 'sponsorship requirement', 'expected salary',
            'current location', 'available start date', 'legal attestations', 'criminal history',
            'EEO/disability/veteran/demographic information', 'background check authorization'
        ]
    },
    'notes': 'MVP defaults from Job Auto Apply PRD_Simplify.md; confirm target regions/industries/roles with user before real search.'
})

write_json_if_missing('profiles/resume_role_mapping.json', {
    'version': 1,
    'mappings': [{
        'profile': 'profiles/profile_en.json',
        'resume_file': 'resumes/pdf/Resume_EN.pdf',
        'autoapply_resume_file': '.autoapply/resumes/Resume_EN.pdf',
        'target_regions': ['Singapore', 'APAC', 'Hong Kong', 'Remote'],
        'industries': ['FinTech', 'Consulting', 'Internet', 'Data Services'],
        'role_categories': [
            'Data Analyst', 'Data Scientist', 'Business Analyst', 'Machine Learning / AI',
            'Quantitative / FinTech', 'Technical Consultant', 'Analytics / Strategy / Operations Analytics'
        ]
    }]
})

write_json_if_missing('profiles/application_private.example.json', {
    'version': 1,
    'warning': 'Example only. Do not store passwords here. Sensitive application answers require user confirmation.',
    'phone_candidates': [],
    'current_location': {'value': '', 'autoFill': False},
    'work_authorization': {'value': '', 'autoFill': False},
    'visa_sponsorship': {'value': '', 'autoFill': False},
    'expected_salary': {'value': '', 'autoFill': False},
    'available_start_date': {'value': '', 'autoFill': False}
})

now = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
for rel, title in [
    ('logs/autofill_runs.md', 'Autofill Runs'),
    ('logs/job_search_runs.md', 'Job Search Runs'),
    ('logs/application_status_checks.md', 'Application Status Checks'),
]:
    p = root / rel
    if not p.exists():
        p.write_text(f'# {title}\n\n- Created: {now}\n- Policy: semi-automated only; never auto-submit.\n\n', encoding='utf-8')
        created.append(rel)

p = root / 'logs/setup_runs.md'
entry = f'\n## {now} PRD MVP structure setup\n\n- Created missing MVP directories/templates.\n- Copied project root Resume.pdf into resume locations non-destructively.\n- Created Excel workbook template if missing.\n'
if p.exists():
    with p.open('a', encoding='utf-8') as f:
        f.write(entry)
else:
    p.write_text('# Setup Runs\n' + entry, encoding='utf-8')
    created.append('logs/setup_runs.md')

scripts = {
    'scripts/convert_resume_profile_en.py': '"""MVP placeholder: convert/refresh English resume profile.\n\nNo-op by default. Future implementation should read resumes/pdf/Resume_EN.pdf\nand update profiles/profile_en.json without inventing facts.\n"""\nfrom pathlib import Path\n\nROOT = Path(__file__).resolve().parents[1]\nprint({"status": "placeholder", "profile": str(ROOT / "profiles" / "profile_en.json")})\n',
    'scripts/update_job_excel.py': '"""MVP placeholder: append/update job-search rows in the master workbook.\n\nMust not delete rows or overwrite user-entered fields without confirmation.\n"""\nfrom pathlib import Path\n\nROOT = Path(__file__).resolve().parents[1]\nprint({"status": "placeholder", "workbook": str(ROOT / "excel" / "岗位搜索与简历投递.xlsx")})\n',
    'scripts/status_check.py': '"""MVP placeholder: check application statuses.\n\nMust not log in automatically or change Excel when no clear status change exists.\n"""\nfrom pathlib import Path\n\nROOT = Path(__file__).resolve().parents[1]\nprint({"status": "placeholder", "workbook": str(ROOT / "excel" / "岗位搜索与简历投递.xlsx")})\n',
}
for rel, content in scripts.items():
    p = root / rel
    if not p.exists():
        p.write_text(content, encoding='utf-8')
        created.append(rel)

xlsx_path = root / 'excel/岗位搜索与简历投递.xlsx'
excel_action = 'exists'
if not xlsx_path.exists():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    schemas = {
        '意向投递公司': ['company_id', 'Company Name', '国家/地区', '城市', '是否目标地区有分公司', '行业', '岗位大类', '匹配简历', 'Career 官网链接', 'ATS 类型', 'ATS Board 链接', 'LinkedIn 公司岗位链接', '信息来源', '最近检查时间', '备注'],
        '开放岗位列表': ['job_id', 'company_id', 'Company Name', 'Job Title', 'Role Category', 'Country/Region', 'City', 'Work Mode', 'Language', 'JD Summary', 'Key Requirements', 'JD URL', 'Application URL', 'Source Platform', 'Match Score', 'Match Rationale', 'Risk / Gaps', 'Discovered At', 'Last Checked At', 'Is Applied', 'User Decision', 'Notes'],
        '已投递记录': ['application_id', 'job_id', 'Company Name', 'Job Title', 'Country/Region', 'City', 'Resume/Profile Used', 'JD URL', 'Application URL', 'Platform', 'Applied At', 'Tool Combination', 'User Submitted Final Application', 'Current Status', 'Status Updated At', 'Status 1', 'Updated At 1', 'Status 2', 'Updated At 2', 'Status 3', 'Updated At 3', 'Status 4', 'Updated At 4', 'Notes'],
    }
    first = True
    for sheet, headers in schemas.items():
        ws = wb.active if first else wb.create_sheet(sheet)
        ws.title = sheet
        first = False
        ws.append(headers)
        ws.append([''] * len(headers))
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}1'
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='1F4E78')
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for idx, h in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = max(14, min(36, len(h) + 4))
        table = Table(
            displayName={'意向投递公司': 'TargetCompanies', '开放岗位列表': 'OpenJobs', '已投递记录': 'AppliedRecords'}[sheet],
            ref=f'A1:{get_column_letter(len(headers))}2'
        )
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)
    wb.properties.title = '岗位搜索与简历投递'
    wb.properties.subject = 'Case A Job Auto Apply MVP tracker'
    wb.save(xlsx_path)
    excel_action = 'created'
    created.append('excel/岗位搜索与简历投递.xlsx')

print(json.dumps({
    'resume_source': str(resume_src),
    'resume_source_sha256': sha(resume_src),
    'copy_results': copy_results,
    'resume_data_action': resume_data_action,
    'excel_action': excel_action,
    'created': created,
}, ensure_ascii=False, indent=2))
